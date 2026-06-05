import os
import base64
import tempfile
import shutil
from io import BytesIO
import re
from celery import shared_task
from docx import Document
from docx.shared import RGBColor
from latest_model.ai_service import get_ai_service

ai_service = get_ai_service()

REQUIRED_DOCUMENTS = (
    ("plan", "Р—Р°СЏРІРєР° РІ РїР»Р°РЅ-РіСЂР°С„РёРє"),
    ("contract", "РџСЂРѕРµРєС‚ РєРѕРЅС‚СЂР°РєС‚Р°"),
    ("ooz", "РћРћР—"),
    ("zapiska", "РџРѕСЏСЃРЅРёС‚РµР»СЊРЅР°СЏ Р·Р°РїРёСЃРєР°"),
    ("onmck", "РћРќРњР¦Рљ"),
    ("obrasheniye", "РћР±СЂР°С‰РµРЅРёРµ Рѕ РїСЂРѕРІРµРґРµРЅРёРё Р·Р°РєСѓРїРєРё"),
)


def build_result_docx_bytes(ai_response: str) -> bytes:
    """
    РЎРѕР±РёСЂР°РµС‚ docx-С„Р°Р№Р» РёР· С‚РµРєСЃС‚РѕРІРѕРіРѕ РѕС‚РІРµС‚Р° РјРѕРґРµР»Рё.

    РџРѕРґРґРµСЂР¶РёРІР°РµС‚ Р±Р°Р·РѕРІРѕРµ С„РѕСЂРјР°С‚РёСЂРѕРІР°РЅРёРµ:
    `<b>...</b>` -> Р¶РёСЂРЅС‹Р№, `<u>...</u>` Рё `<ins>...</ins>` -> РїРѕРґС‡С‘СЂРєРёРІР°РЅРёРµ,
    `<ok>...</ok>` -> Р·РµР»С‘РЅС‹Р№ С‚РµРєСЃС‚, `<warn>...</warn>` -> РѕСЂР°РЅР¶РµРІС‹Р№ С‚РµРєСЃС‚,
    `<error>...</error>` -> РєСЂР°СЃРЅС‹Р№ С‚РµРєСЃС‚.
    РђР±Р·Р°С†С‹ СЃРѕР·РґР°СЋС‚СЃСЏ РїРѕ РїСѓСЃС‚С‹Рј СЃС‚СЂРѕРєР°Рј.
    """
    document = Document()
    document.add_heading('Р РµР·СѓР»СЊС‚Р°С‚ РїСЂРѕРІРµСЂРєРё РґРѕРєСѓРјРµРЅС‚РѕРІ', level=1)

    clean_response = (ai_response or '').replace('\r\n', '\n')
    blocks = [block.strip() for block in clean_response.split('\n\n') if block.strip()]
    tag_pattern = re.compile(r"</?(?:b|u|ins|ok|warn|error)>", re.IGNORECASE)

    for block in blocks:
        paragraph = document.add_paragraph()
        bold_active = False
        underline_active = False
        ok_active = False
        warn_active = False
        error_active = False
        cursor = 0

        for match in tag_pattern.finditer(block):
            if match.start() > cursor:
                run = paragraph.add_run(block[cursor:match.start()])
                run.bold = bold_active
                run.underline = underline_active
                if ok_active:
                    run.font.color.rgb = RGBColor(0x19, 0x87, 0x54)
                elif warn_active:
                    run.font.color.rgb = RGBColor(0xFD, 0x7E, 0x14)
                elif error_active:
                    run.font.color.rgb = RGBColor(0xDC, 0x35, 0x45)

            tag = match.group(0).lower()
            if tag == "<b>":
                bold_active = True
            elif tag == "</b>":
                bold_active = False
            elif tag in ("<u>", "<ins>"):
                underline_active = True
            elif tag in ("</u>", "</ins>"):
                underline_active = False
            elif tag == "<ok>":
                ok_active = True
            elif tag == "</ok>":
                ok_active = False
            elif tag == "<warn>":
                warn_active = True
            elif tag == "</warn>":
                warn_active = False
            elif tag == "<error>":
                error_active = True
            elif tag == "</error>":
                error_active = False

            cursor = match.end()

        if cursor < len(block):
            run = paragraph.add_run(block[cursor:])
            run.bold = bold_active
            run.underline = underline_active
            if ok_active:
                run.font.color.rgb = RGBColor(0x19, 0x87, 0x54)
            elif warn_active:
                run.font.color.rgb = RGBColor(0xFD, 0x7E, 0x14)
            elif error_active:
                run.font.color.rgb = RGBColor(0xDC, 0x35, 0x45)

    buffer = BytesIO()
    document.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


@shared_task(bind=True, name='rag_worker.process_document_query')
def process_document_query(self, documents):
    try:
        if not isinstance(documents, list):
            raise ValueError("Expected a list of uploaded documents.")

        docs_by_key = {}
        for document in documents:
            doc_key = document.get('key')
            file_name = document.get('name')
            file_content_b64 = document.get('content_b64')

            if not doc_key or not file_name or not file_content_b64:
                raise ValueError("Each uploaded document must contain key, name and content_b64.")

            docs_by_key[doc_key] = {
                'name': file_name,
                'content': base64.b64decode(file_content_b64),
            }

        missing_docs = [label for key, label in REQUIRED_DOCUMENTS if key not in docs_by_key]
        if missing_docs:
            raise ValueError(
                f"Missing required documents: {', '.join(missing_docs)}."
            )

        temp_dir = tempfile.mkdtemp()
        try:
            doc_paths = {}
            for key, _label in REQUIRED_DOCUMENTS:
                file_name = os.path.basename(docs_by_key[key]['name'])
                temp_file_path = os.path.join(temp_dir, f"{key}_{file_name}")

                with open(temp_file_path, 'wb') as f:
                    f.write(docs_by_key[key]['content'])

                doc_paths[key] = temp_file_path

            result = ai_service.process_query(
                plan_path=doc_paths['plan'],
                contract_path=doc_paths['contract'],
                ooz_path=doc_paths['ooz'],
                zapiska_path=doc_paths['zapiska'],
                ONMCK_path=doc_paths['onmck'],
                Obrasheniye_path=doc_paths['obrasheniye'],
            )
            result_file_bytes = build_result_docx_bytes(result['ai_response'])

            return {
                'ai_response': result['ai_response'],
                'result_file_b64': base64.b64encode(result_file_bytes).decode('utf-8'),
                'result_file_name': 'analysis_result.docx',
                'documents': [
                    {
                        'key': key,
                        'label': label,
                        'name': docs_by_key[key]['name'],
                    }
                    for key, label in REQUIRED_DOCUMENTS
                ],
                'status': 'completed'
            }
        finally:
            if os.path.exists(temp_dir):
                shutil.rmtree(temp_dir)
    except Exception as e:
        error_msg = str(e)
        print(f"Error processing documents: {error_msg}")
        raise Exception(error_msg)

